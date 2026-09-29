"""Build the four-page IEEESTEC-style JazzDialog conference manuscript."""

from pathlib import Path
import re

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "rad"
FIG = OUT / "figure" / "ieeestec_good_bad_revidiran.png"
RANK_FIG = OUT / "figure" / "ieeestec_ranking_final_draft.png"
DTW_FIG = OUT / "figure" / "ieeestec_dtw_overlap_final_draft.png"


def set_font(style, name="Times New Roman", size=9, bold=False):
    style.font.name = name
    style.font.size = Pt(size)
    style.font.bold = bold
    style.font.color.rgb = RGBColor(0, 0, 0)
    fonts = style.element.get_or_add_rPr().get_or_add_rFonts()
    for attr in ("ascii", "hAnsi", "eastAsia", "cs"):
        fonts.set(qn(f"w:{attr}"), name)
    lang = OxmlElement("w:lang")
    lang.set(qn("w:val"), "sr-Latn-RS")
    style.element.get_or_add_rPr().append(lang)


def two_columns(section):
    cols = section._sectPr.xpath("./w:cols")[0]
    cols.set(qn("w:num"), "2")
    cols.set(qn("w:space"), "300")
    cols.set(qn("w:equalWidth"), "1")


def one_column(section):
    cols = section._sectPr.xpath("./w:cols")[0]
    cols.set(qn("w:num"), "1")


def compact_example_figure():
    """Create a compact, data-preserving view of the two reviewed examples."""
    import matplotlib.pyplot as plt
    good_call = [71, 72, 74, 67, 67, 67, 67, 68, 67, 66, 67]
    good_response = [69, 70, 72, 73, 64, 64, 65, 67, 68]
    bad_call = [73, 72, 70, 72, 73, 72, 70]
    bad_response = [73, 73, 72, 70, 72, 70]
    examples = [
        ("(a) I Fall in Love Too Easily — prihvaćen\nDTW = 0,464",
         good_call, good_response),
        ("(b) Blues for Blanche — odbijen\nDTW = 0,125",
         bad_call, bad_response),
    ]
    fig, axes = plt.subplots(1, 2, figsize=(6.5, 2.55), sharey=True)
    for ax, (title, call, response) in zip(axes, examples):
        ax.plot(range(1, len(call) + 1), call, "o-", color="#2878b5",
                linewidth=1.25, markersize=3, label="Call")
        ax.plot(range(1, len(response) + 1), response, "s-", color="#d67929",
                linewidth=1.25, markersize=3, label="Response")
        ax.set_title(title, fontsize=8, pad=5)
        ax.set_ylabel("MIDI visina", fontsize=7)
        ax.grid(alpha=.25, linewidth=.4)
        ax.tick_params(labelsize=6)
        ax.legend(loc="upper right", fontsize=6, frameon=False)
    for ax in axes:
        ax.set_xlabel("Redni broj note", fontsize=7)
    fig.tight_layout(pad=.65)
    FIG.parent.mkdir(parents=True, exist_ok=True)
    fig.savefig(FIG, dpi=320, bbox_inches="tight")
    plt.close(fig)


def compact_ranking_figure():
    """Plot only the ranking results needed in the conference manuscript."""
    import matplotlib.pyplot as plt
    names = ["DTW + incipit", "Rang sa više\nosobina", "Blaga kazna\nintervala"]
    all_scores = [42.78, 56.11, 53.89]
    mlu_scores = [44.67, 54.67, 56.00]
    pos = range(len(names))
    fig, ax = plt.subplots(figsize=(5.8, 6.0))
    ax.barh([p - .18 for p in pos], all_scores, height=.34, color="#356b88",
            label="Svi automatski kandidati (n=299)")
    ax.barh([p + .18 for p in pos], mlu_scores, height=.34, color="#cf8346",
            label="MLU podskup (n=238)")
    ax.set_yticks(list(pos), names, fontsize=8)
    ax.invert_yaxis()
    ax.set_xlim(0, 65)
    ax.set_xlabel("DA među prvih 20% kandidata za pregled (%)", fontsize=8)
    ax.tick_params(labelsize=8)
    ax.grid(axis="x", alpha=.25, linewidth=.4)
    ax.legend(fontsize=7, loc="lower right")
    fig.tight_layout(pad=.45)
    fig.savefig(RANK_FIG, dpi=320, bbox_inches="tight")
    plt.close(fig)


def para(doc, text="", style=None, align=WD_ALIGN_PARAGRAPH.JUSTIFY):
    p = doc.add_paragraph(style=style)
    p.alignment = align
    p.paragraph_format.space_after = Pt(3)
    p.paragraph_format.line_spacing = 1.05
    p.add_run(text)
    return p


def heading(doc, text):
    p = para(doc, text, "Heading 1", WD_ALIGN_PARAGRAPH.LEFT)
    p.paragraph_format.space_before = Pt(5)
    p.paragraph_format.space_after = Pt(2)
    return p


def subheading(doc, text):
    p = para(doc, text, "Heading 2", WD_ALIGN_PARAGRAPH.LEFT)
    p.paragraph_format.space_before = Pt(3)
    p.paragraph_format.space_after = Pt(1)
    return p


def picture(doc, path, caption):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(2)
    p.paragraph_format.space_after = Pt(1)
    p.add_run().add_picture(str(path), width=Cm(7.55))
    c = para(doc, caption, "Caption", WD_ALIGN_PARAGRAPH.JUSTIFY)
    c.paragraph_format.space_after = Pt(3)


def wide_picture(doc, path, caption, width=15.6):
    before = doc.add_section(WD_SECTION.CONTINUOUS)
    before.page_width, before.page_height = Cm(21), Cm(29.7)
    before.top_margin, before.bottom_margin = Cm(1.905), Cm(2.54)
    before.left_margin, before.right_margin = Cm(1.6), Cm(1.6)
    one_column(before)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(3)
    p.paragraph_format.space_after = Pt(1)
    p.add_run().add_picture(str(path), width=Cm(width))
    c = para(doc, caption, "Caption", WD_ALIGN_PARAGRAPH.JUSTIFY)
    c.paragraph_format.space_after = Pt(3)
    after = doc.add_section(WD_SECTION.CONTINUOUS)
    after.page_width, after.page_height = Cm(21), Cm(29.7)
    after.top_margin, after.bottom_margin = Cm(1.905), Cm(2.54)
    after.left_margin, after.right_margin = Cm(1.6), Cm(1.6)
    two_columns(after)


def dataset_summary(doc):
    """Compact summary of the final dataset, using verified corpus counts."""
    table = doc.add_table(rows=3, cols=2)
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    values = [
        ("114", "potvrđenih parova"),
        ("76", "solaža / 73 naslova"),
        ("42", "izvođača; 26 unutar i 88 preko granice fraze"),
    ]
    for row, (value, label) in zip(table.rows, values):
        row.cells[0].text = value
        row.cells[1].text = label
        for run in row.cells[0].paragraphs[0].runs:
            run.bold = True
            run.font.size = Pt(10)
        for cell in row.cells:
            for paragraph in cell.paragraphs:
                paragraph.paragraph_format.space_after = Pt(0)
                paragraph.paragraph_format.line_spacing = 1.0
                for run in paragraph.runs:
                    run.font.name = "Times New Roman"
                    run.font.size = Pt(8)
    doc.add_paragraph().paragraph_format.space_after = Pt(1)


def build():
    if not FIG.exists():
        compact_example_figure()
    if not RANK_FIG.exists():
        compact_ranking_figure()
    doc = Document()
    first = doc.sections[0]
    first.page_width, first.page_height = Cm(21), Cm(29.7)
    first.top_margin, first.bottom_margin = Cm(1.905), Cm(2.54)
    first.left_margin, first.right_margin = Cm(1.6), Cm(1.6)

    set_font(doc.styles["Normal"], size=10)
    set_font(doc.styles["Title"], size=16, bold=True)
    set_font(doc.styles["Heading 1"], size=10, bold=True)
    set_font(doc.styles["Heading 2"], size=9, bold=True)
    set_font(doc.styles["Caption"], size=8)
    doc.styles["Normal"].paragraph_format.space_after = Pt(0)
    doc.styles["Normal"].paragraph_format.line_spacing = 1.0

    header = first.header.paragraphs[0]
    header.text = "IEEESTEC – Studentska konferencija, Niš, 2026."
    header.alignment = WD_ALIGN_PARAGRAPH.LEFT
    for run in header.runs:
        run.font.name = "Times New Roman"
        run.font.size = Pt(8)

    title = doc.add_paragraph("Izdvajanje call-and-response parova iz jazz improvizacija: JazzDialog baza", "Title")
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.paragraph_format.space_after = Pt(5)

    authors = doc.add_table(rows=1, cols=1)
    authors.alignment = WD_TABLE_ALIGNMENT.CENTER
    cell = authors.cell(0, 0)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    p = cell.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(0)
    r = p.add_run("Nikolina Zdravković")
    r.bold = True
    r.font.size = Pt(9)
    p = cell.add_paragraph("Istraživačka stanica Petnica, Valjevo, Srbija")
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(0)
    p.runs[0].font.size = Pt(8)
    p = cell.add_paragraph("Mentorka: Maja Milović")
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(2)
    p.runs[0].font.size = Pt(8)
    for border in cell._tc.tcPr.xpath("./w:tcBorders/*"):
        border.getparent().remove(border)

    body = doc.add_section(WD_SECTION.CONTINUOUS)
    body.page_width, body.page_height = Cm(21), Cm(29.7)
    body.top_margin, body.bottom_margin = Cm(1.905), Cm(2.54)
    body.left_margin, body.right_margin = Cm(1.6), Cm(1.6)
    two_columns(body)

    p = para(doc, "Sažetak — U radu je prikazan postupak izdvajanja call-and-response parova iz transkribovanih jazz solaža i nastanak baze JazzDialog. Kandidati su dobijani iz melodijskih fraza, poređeni kroz više muzički motivisanih osobina, a zatim rangirani za ručni pregled. Analiza pokazuje da mala melodijska udaljenost nije dovoljna potvrda call-and-response odnosa: odbijeni kandidat može imati niži DTW skor od prihvaćenog. Konačni postupak zato kombinuje računarsko izdvajanje i rangiranje sa slušanjem i ručnom potvrdom. JazzDialog sadrži 114 potvrđenih zapisa iz 76 solaža, 73 naslova i 42 izvođača, sa preciznim granicama segmenata i MIDI isečcima.")
    p.runs[0].bold = True
    p = para(doc, "Ključne reči — jazz improvizacija, call-and-response, analiza melodije, DTW, MIDI, baza podataka")
    p.runs[0].italic = True

    heading(doc, "I. UVOD")
    para(doc, "Call-and-response je odnos dve uzastopne muzičke celine u kome druga celina odgovara, menja ili razvija ideju prve [5, 6]. U jazz improvizaciji odgovor ne mora biti doslovno ponavljanje: može preuzeti samo deo motiva, promeniti ritam ili početnu visinu, produžiti ideju, promeniti završetak ili odgovoriti kontrastnom frazom. Improvizacija zato dozvoljava veliku varijabilnost načina na koji se muzička ideja razvija.")
    para(doc, "Ta sloboda je računarski problem. Pravilo koje bi dobro opisalo jedan validan par može odbaciti drugi, dok prosto ponavljanje nekoliko tonova može izgledati vrlo slično bez funkcije pitanja i odgovora. Zbog toga sličnost dve uzastopne fraze nije dovoljna pretpostavka za call-and-response, već je potrebno proveriti više osobina i na kraju preslušati kandidat.")
    para(doc, "Cilj rada bio je da se ispita kako iz transkribovanih jazz improvizacija računarski izdvojiti kandidate za call-and-response i u kojoj meri se muzička procena tog odnosa može opisati numeričkim osobinama melodijskih fraza. Konačni rezultat je JazzDialog, skup ručno potvrđenih parova sa vezom ka izvornim transkripcijama i MIDI isečcima.")

    heading(doc, "II. METOD")
    subheading(doc, "A. Od sola do kandidata")
    para(doc, "Korišćene su monofone MIDI transkripcije stvarnih jazz standard solaža i prateće anotacije sistema Jazzomat / Weimar Jazz Database [1]. Postupak je imao sledeći tok: solo se deli na početne muzičke celine, iz njih se formiraju mogući call/response parovi, za svaki par se računaju osobine, kandidati se rangiraju, preslušavaju i označavaju kao prihvaćeni ili odbijeni. Tek prihvaćeni parovi ulaze u JazzDialog.")
    para(doc, "Pre poređenja je bilo potrebno odlučiti koje delove sola međusobno posmatrati. Prvo je testirano deljenje po pauzama, ali pauza nije pouzdana granica muzičke ideje. Klizni prozori kroz ceo solo davali su veliki broj lokalnih sličnosti koje nisu nužno bile povezane muzičke celine. Zato su zvanično anotirane granice fraza korišćene kao stabilniji početni okvir za generisanje kandidata. One nisu definicija call-and-response odnosa: konačna baza sadrži i potvrđene parove koji prelaze granicu jedne anotirane fraze.")
    para(doc, "U osnovnoj pretrazi fraza se deli na dva uzastopna dela: prvi je call, a drugi response. Kandidati sa premalo nota, prekratkim trajanjem ili izrazito neujednačenim dužinama su odbačeni. U kasnijoj verziji response je imao između polovine i dvostruke dužine call-a, a oba dela su trajala najmanje jednu sekundu. Ova ograničenja samo uklanjaju trivijalne slučajeve; ne predstavljaju definiciju odnosa.")

    subheading(doc, "B. Numeričke osobine i rangiranje")
    para(doc, "Pošto response ne mora biti kopija call-a, analizirane su različite vrste sličnosti. DTW ispituje da li dve melodijske linije imaju sličan tok čak i kada nemaju isti broj nota [2]. Transponovana, odnosno shape predstava, zanemaruje početnu visinu da bi proverila oblik melodije. Incipit proverava da li response na početku preuzima ideju call-a, intervali obrazac kretanja između tonova, kontura opšti smer melodije, a relativni ritam odnose razmaka između nota.")
    para(doc, "DTW je računat nad MIDI visinama uz optimalni put kroz matricu razlika; manja vrednost označava bliže melodijsko poravnanje. U brzoj frazi incipit je proširivan do približno 0,75 s, jer tri kratke note nisu dovoljan signal. Kandidati su zatim rangirani kombinovanjem DTW-a, incipita, odnosa dužina, zastupljenosti istih tonova, raznovrsnosti visina i približnog ponavljanja intervala.")
    subheading(doc, "C. Potvrda")
    para(doc, "Računarski postupak nije donosio konačnu odluku da je kandidat call-and-response. Njegova uloga bila je da izdvoji i rangira kandidate, nakon čega je svaki preslušan i ručno označen kao prihvaćen ili odbijen na osnovu muzičke procene odnosa. Konačnu oznaku davala je autorka rada.")

    heading(doc, "III. REZULTATI I DISKUSIJA")
    para(doc, "Prvi dokaz ograničenja DTW-a je konkretan par primera na slici 1. Prihvaćeni kandidat iz sola „I Fall in Love Too Easily“ ima DTW skor 0,464, dok odbijeni kandidat iz „Blues for Blanche“ ima znatno manji skor, 0,125. Kako manji skor znači veću melodijsku sličnost, prost kriterijum bi pogrešno favorizovao odbijeni primer. On sadrži kratko, mehaničko ponavljanje tonova koje algoritam lako poravnava, ali pri slušanju ne funkcioniše kao jasan odgovor. Prihvaćeni response nije kopija call-a, već razvija prethodnu muzičku ideju. Melodijska sličnost zato nije isto što i call-and-response funkcija.")
    p = para(doc, "Primeri sa slike 1: (a) I Fall in Love Too Easily, melid 70, fraza 5; (b) Blues for Blanche, melid 2, fraza 39.", align=WD_ALIGN_PARAGRAPH.CENTER)
    for run in p.runs:
        run.bold = True
        run.font.size = Pt(7.6)
    p.paragraph_format.space_after = Pt(0)
    wide_picture(doc, FIG, "Slika 1. Melodijske konture call-a i response-a za (a) I Fall in Love Too Easily i (b) Blues for Blanche. Horizontalna osa je redni broj note unutar segmenta, a vertikalna osa MIDI visina. Odbijeni primer ima znatno niži DTW skor, što pokazuje da skor sam nije odluka o muzičkoj funkciji.", width=11.0)
    para(doc, "Dva primera mogu biti izolovan slučaj, pa je zatim provereno da li se DTW skorovi razdvajaju na većem skupu od 299 ručno ocenjenih kandidata: 101 prihvaćenom i 198 odbijenih. Ako bi DTW bio pouzdan samostalan kriterijum, prihvaćeni kandidati bi uglavnom imali niže skorove. Slika 2 pokazuje veliko preklapanje raspodela; zato ne postoji jednostavan DTW prag koji pouzdano odvaja call-and-response od ostalih kandidata. Figura prikazuje pojedinačne kandidate i njihovu raspodelu po dve ručne oznake.")
    wide_picture(doc, DTW_FIG, "Slika 2. Raspodele normalizovanog globalnog DTW skora za ručno prihvaćene i odbijene automatske kandidate.")
    para(doc, "Pošto nijedna mera nije mogla sama da potvrdi odnos, sledeći cilj nije bila automatska klasifikacija već bolji redosled za ručni pregled: dobri kandidati treba da budu što bliže vrhu liste, kako bi se preslušalo manje loših primera. Metrika na slici 3 je procenat kandidata koji su među prvih 20% rangirane liste i nakon slušanja dobijaju oznaku DA; to nije accuracy klasifikatora. Na zamrznutom razvojnom skupu, uz petostruku podelu po solima i tri semena, DTW sa incipitom daje 42,78%, rang sa više osobina 56,11%, a blaga kazna intervalskih obrazaca 53,89%. U trećem postupku obrazac koji je čest među razvojnim NE kandidatima samo snižava prioritet kandidata, bez automatskog odbacivanja. Više osobina zato poboljšava kvalitet liste, ali ne uklanja potrebu za slušanjem.")
    wide_picture(doc, RANK_FIG, "Slika 3. Udeo ručno prihvaćenih kandidata u prvih 20% rangirane liste. Stubovi upoređuju sve automatske kandidate i MLU podskup; prikaz je sažet na tri najinformativnija postupka.")
    para(doc, "Na ranom označenom skupu ispitano je više varijanti melodijske udaljenosti. Najbolja je dostigla AUC 0,603, dok je jednostavna ritmička mera bila blizu slučajnog razdvajanja, AUC 0,501. Pojedinačne mere ritma i konture zato nisu dale stabilnu autonomnu odluku.")
    para(doc, "Kontura je zato posmatrana kao dodatna, a ne kao samostalna odluka. Ideju je podstakla statistička karakterizacija melodijskih kontura [4], ali se postupak ne može direktno preneti: taj rad izdvaja vodeću melodiju iz polifonog zvuka, dok su ovde već dostupne monofone MIDI note i treba proceniti odnos dva segmenta. Sažeta kontura i njena kombinacija sa ritmom nisu dale stabilan napredak.")
    subheading(doc, "JazzDialog baza")
    para(doc, "Konačni skup JazzDialog sadrži 114 ručno potvrđenih call-and-response zapisa iz 76 solaža, 73 naslova i 42 izvođača. Svaki zapis čuva identifikator sola, izvođača, naslov, granice i vremena call/response delova, nizove MIDI visina i putanju do MIDI isečka. Call delovi imaju medijanu od 9, a response delovi medijanu od 10 nota. Dvadeset šest anotacija u celosti je unutar jedne zvanične fraze, dok 88 prelazi njenu granicu; fraze su dakle koristan okvir za pretragu, ali ne i potpuna definicija muzičke ideje.")
    dataset_summary(doc)
    picture(doc, OUT / "figure" / "figure_1_annotation_scope.png", "Slika 4. Obuhvat zvaničnih granica fraze u potvrdama JazzDialog baze. Granice fraze su koristan početni okvir, ali većina potvrđenih parova prelazi granicu jedne fraze.")
    para(doc, "Za razliku od skupova namenjenih generisanju muzičkog odgovora [3], JazzDialog polazi od stvarnih improvizacija i dokumentuje postupak pronalaženja i ručne potvrde odnosa koji se već dogodio. Njegov doprinos je proverljiva kolekcija i negativan rezultat važan za dalji rad: melodijska sličnost hvata deo odnosa, ali ne obuhvata njegovu muzičku funkciju.")

    heading(doc, "IV. OGRANIČENJA I DALJI RAD")
    para(doc, "Najvažnije praktično ograničenje jeste da ispitane numeričke osobine i modeli rangiranja nisu dostigli pouzdanost dovoljnu da se konačna odluka prepusti algoritmu. Zato je svaki kandidat morao biti preslušan i ručno potvrđen. Call-and-response je delimično subjektivan pojam, a u ovoj verziji konačnu oznaku davala je jedna anotatorka.")
    para(doc, "Korišćene su monofone MIDI transkripcije, pa nisu predstavljeni artikulacija, dinamika, boja tona, harmonijska pratnja ni puna interakcija sa ansamblom. Razvojne oznake su korišćene i pri oblikovanju kriterijuma i pri retrospektivnom poređenju rangiranja; podela po solima smanjuje curenje srodnih primera, ali nije potpuno nezavisna završna evaluacija.")
    para(doc, "Sledeći korak je ljudska evaluacija sa nezavisnim slušaocima: obrazovani ili iskusni jazz muzičari i slušaoci sa manjim jazz iskustvom ocenjivali bi iste primere. Time bi se ispitalo koliko se njihove procene međusobno slažu, koliko se slažu sa postojećim oznakama i da li jazz obrazovanje utiče na percepciju odnosa.")
    para(doc, "JazzDialog je početna verzija baze i treba je proširiti većim brojem solaža, potvrđenih parova i anotatora. Veći skup bi omogućio pouzdaniju analizu i modele koji bi učili obrasce muzičkog odgovora iz stvarnih improvizacija. Dugoročno, baza može biti osnova za modele koji generišu response oslanjajući se na zabeležene obrasce razvoja i transformacije muzičke ideje; to nije realizovan rezultat ovog rada.")

    heading(doc, "V. ZAKLJUČAK")
    para(doc, "Analizirane su stvarne jazz improvizacije da bi se izdvojili kandidati za call-and-response i ispitano je više načina njihovog numeričkog opisivanja. Melodijska sličnost, uključujući DTW, nije bila dovoljna za potvrdu odnosa, dok kombinovanje osobina može poboljšati rangiranje za ručni pregled. Hibridnim postupkom izdvajanja, rangiranja i preslušavanja nastala je JazzDialog baza sa 114 potvrđenih parova. Baza je osnova za proširenje anotacija, ljudsku evaluaciju i buduće modele jazz call-and-response improvizacije.")

    heading(doc, "LITERATURA")
    references = [
        "[1] M. Pfleiderer, K. Frieler, J. Abeßer, W. G. Zaddach i B. Burkhart, ur., Inside the Jazzomat: New Perspectives for Jazz Research. Mainz: Schott Campus, 2017.",
        "[2] M. Müller, Fundamentals of Music Processing: Audio, Analysis, Algorithms, Applications. Cham: Springer, 2015, doi: 10.1007/978-3-319-21945-5.",
        "[3] Y. Hu, Z. Wang, R. Liu, Y. Liang i Y. Zhang, “Responding to the Call: Exploring Automatic Music Composition Using a Knowledge-Enhanced Model,” Proc. AAAI Conf. Artif. Intell., vol. 38, no. 1, pp. 521–529, 2024, doi: 10.1609/aaai.v38i1.27807.",
        "[4] J. Salamon, G. Peeters i A. Röbel, “Statistical Characterisation of Melodic Pitch Contours and Its Application for Melody Extraction,” Proc. 13th Int. Soc. Music Information Retrieval Conf., pp. 187–192, 2012.",
        "[5] P. F. Berliner, Thinking in Jazz: The Infinite Art of Improvisation. Chicago: University of Chicago Press, 1994.",
        "[6] I. Monson, Saying Something: Jazz Improvisation and Interaction. Chicago: University of Chicago Press, 1996.",
    ]
    for ref in references:
        p = para(doc, ref, align=WD_ALIGN_PARAGRAPH.LEFT)
        p.paragraph_format.left_indent = Cm(.18)
        p.paragraph_format.first_line_indent = Cm(-.18)

    doc.core_properties.title = "Izdvajanje call-and-response parova iz jazz improvizacija: JazzDialog baza"
    doc.core_properties.author = "Nikolina Zdravković"
    doc.core_properties.subject = "IEEESTEC konferencijski rad"
    output = OUT / "ieeestec_jazzdialog_final_draft.docx"
    doc.save(output)
    print(output)


if __name__ == "__main__":
    build()
