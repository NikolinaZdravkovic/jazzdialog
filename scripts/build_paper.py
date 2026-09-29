"""Build the editable manuscript from DRAFT_RADA.md (requires python-docx).

Figures: python scripts/make_paper_figures.py. PDF is exported from the DOCX
using Word/LibreOffice and must be visually checked before distribution.
"""
import re
from pathlib import Path

from docx import Document
from docx.shared import Cm, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT = Path(__file__).resolve().parents[1]


def keep(paragraph, next_one=False):
    paragraph.paragraph_format.keep_together = True
    paragraph.paragraph_format.keep_with_next = next_one


def build():
    doc = Document()
    sec = doc.sections[0]
    sec.page_width, sec.page_height = Cm(21), Cm(29.7)
    sec.top_margin, sec.bottom_margin = Cm(2), Cm(2)
    sec.left_margin, sec.right_margin = Cm(2.3), Cm(2.3)
    for name in ['Normal', 'Title', 'Heading 1', 'Heading 2', 'Caption']:
        style = doc.styles[name]
        style.font.name = 'Times New Roman'
        style.font.color.rgb = RGBColor(0, 0, 0)
        fonts = style.element.get_or_add_rPr().get_or_add_rFonts()
        for attr in list(fonts.attrib):
            if 'Theme' in attr:
                del fonts.attrib[attr]
        for attr in ['ascii', 'hAnsi', 'eastAsia', 'cs']:
            fonts.set(qn('w:'+attr), 'Times New Roman')
        for border in style.element.xpath('.//w:pBdr'):
            border.getparent().remove(border)
        lang = OxmlElement('w:lang'); lang.set(qn('w:val'), 'sr-Latn-RS')
        style.element.get_or_add_rPr().append(lang)
    normal = doc.styles['Normal']
    normal.font.size = Pt(11)
    normal.paragraph_format.line_spacing = 1.04
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.widow_control = True
    for name, size in [('Title', 18), ('Heading 1', 13), ('Heading 2', 11.5)]:
        doc.styles[name].font.size = Pt(size)
        doc.styles[name].font.bold = True
        doc.styles[name].paragraph_format.space_before = Pt(12 if name != 'Title' else 0)
        doc.styles[name].paragraph_format.space_after = Pt(6)
    doc.styles['Caption'].font.size = Pt(9)
    doc.styles['Caption'].font.bold = False
    doc.styles['Caption'].paragraph_format.space_after = Pt(8)
    header = sec.header.paragraphs[0]
    header.text = 'JazzDialog | Nikolina Zdravković'
    header.style = doc.styles['Caption']
    footer = sec.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = footer.add_run()
    field = OxmlElement('w:fldSimple'); field.set(qn('w:instr'), 'PAGE')
    run._r.addnext(field)
    doc.core_properties.title = 'JazzDialog: izgradnja i analiza zbirke call-and-response parova u jazz solažima'
    doc.core_properties.author = 'Nikolina Zdravković'
    doc.core_properties.subject = 'Naučni rad za pregled mentorke'
    blocks = (ROOT/'DRAFT_RADA.md').read_text(encoding='utf-8').strip().split('\n\n')
    bibliography = False
    for block in blocks:
        if block.startswith('# '):
            doc.add_paragraph(block[2:], 'Title')
        elif block.startswith('### '):
            doc.add_paragraph(block[4:], 'Heading 2')
        elif block.startswith('## '):
            bibliography = block == '## Literatura'
            doc.add_paragraph(block[3:], 'Heading 1')
        elif block.startswith('!['):
            path = re.search(r'\]\((.*?)\)', block).group(1)
            p = doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.add_run().add_picture(str(ROOT/path), width=Cm(13.7))
            keep(p, True)
        elif block.startswith(('D(x,y) =', 'S(x,y) =')):
            number = 1 if block.startswith('D(') else 2
            p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.add_run().add_picture(str(ROOT/f'rad/figure/equation_{number}.png'), width=Cm(12.7))
            keep(p)
        elif block.startswith('|'):
            rows = [line.strip('|').split('|') for line in block.splitlines()]
            rows = [row for row in rows if not all(re.fullmatch(r'[\s:-]+', cell) for cell in row)]
            table = doc.add_table(rows=0, cols=len(rows[0]))
            table.style = 'Light Shading Accent 1'
            table.autofit = False
            for column, width in zip(table.columns, [9.4, 3.5, 3.5]):
                column.width = Cm(width)
            for i, row in enumerate(rows):
                cells = table.add_row().cells
                for j, (cell, value) in enumerate(zip(cells, row)):
                    cell.width = Cm(9.4 if j == 0 else 3.5)
                    cell.text = value.strip()
                    for p in cell.paragraphs:
                        p.paragraph_format.space_after = Pt(5)
                        p.paragraph_format.space_before = Pt(5)
                        if j: p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
                        keep(p, i < len(rows)-1)
                        for r in p.runs:
                            r.font.size = Pt(10); r.bold = i == 0
                tr_pr = table.rows[-1]._tr.get_or_add_trPr()
                tr_pr.append(OxmlElement('w:cantSplit'))
                if i == 0: tr_pr.append(OxmlElement('w:tblHeader'))
            doc.add_paragraph().paragraph_format.space_after = Pt(1)
        elif re.match(r'^(Slika|Tabela) \d+\.', block):
            p = doc.add_paragraph(block, 'Caption')
            keep(p, block.startswith('Tabela '))
        else:
            p = doc.add_paragraph(block)
            p.alignment = WD_ALIGN_PARAGRAPH.LEFT if bibliography else WD_ALIGN_PARAGRAPH.JUSTIFY
            if block.startswith(('Nikolina ', 'Mentorka:')):
                p.alignment = WD_ALIGN_PARAGRAPH.LEFT
                keep(p, True)
    output = ROOT/'rad/jazzdialog_rad.docx'
    output.parent.mkdir(parents=True, exist_ok=True)
    doc.save(output)
    print(output)


if __name__ == '__main__':
    build()
