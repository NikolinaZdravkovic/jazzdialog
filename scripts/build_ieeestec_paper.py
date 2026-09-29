"""Build the four-page IEEESTEC-style JazzDialog conference manuscript."""

from pathlib import Path
import re

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "rad"
FIG = OUT / "figure" / "ieeestec_good_bad.png"
RANK_FIG = OUT / "figure" / "ieeestec_ranking.png"


def set_font(style, name="Times New Roman", size=9, bold=False):
    style.font.name = name
    style.font.size = Pt(size)
    style.font.bold = bold
    style.font.color.rgb = RGBColor(0, 0, 0)
    fonts = style.element.get_or_add_rPr().get_or_add_rFonts()
    for attr in ("ascii", "hAnsi", "eastAsia", "cs"):
        fonts.set(qn(f"w:{attr}"), name)
    lang = OxmlElement("w:lang")
    lang.set(qn("w:val"), "sr-Latn-RS")
    style.element.get_or_add_rPr().append(lang)


def two_columns(section):
    cols = section._sectPr.xpath("./w:cols")[0]
    cols.set(qn("w:num"), "2")
    cols.set(qn("w:space"), "300")
    cols.set(qn("w:equalWidth"), "1")


def one_column(section):
    cols = section._sectPr.xpath("./w:cols")[0]
    cols.set(qn("w:num"), "1")


def compact_example_figure():
    """Create a compact, data-preserving view of the two reviewed examples."""
    import matplotlib.pyplot as plt
    good_call = [71, 72, 74, 67, 67, 67, 67, 68, 67, 66, 67]
    good_response = [69, 70, 72, 73, 64, 64, 65, 67, 68]
    bad_call = [73, 72, 70, 72, 73, 72, 70]
    bad_response = [73, 73, 72, 70, 72, 70]
    examples = [
        ("(a) Prihvaćen: I Fall in Love Too Easily (DTW = 0,464)",
         good_call, good_response),
        ("(b) Odbijen: Blues for Blanche (DTW = 0,125)",
         bad_call, bad_response),
    ]
    fig, axes = plt.subplots(2, 1, figsize=(3.4, 4.15), sharey=True)
    for ax, (title, call, response) in zip(axes, examples):
        ax.plot(range(1, len(call) + 1), call, "o-", color="#2878b5",
                linewidth=1.25, markersize=3, label="Call")
        ax.plot(range(1, len(response) + 1), response, "s-", color="#d67929",
                linewidth=1.25, markersize=3, label="Response")
        ax.set_title(title, fontsize=7, pad=3)
        ax.set_ylabel("MIDI visina", fontsize=7)
        ax.grid(alpha=.25, linewidth=.4)
        ax.tick_params(labelsize=6)
        ax.legend(loc="upper right", fontsize=6, frameon=False)
    axes[-1].set_xlabel("Redni broj note unutar segmenta", fontsize=7)
    fig.tight_layout(pad=.55)
    FIG.parent.mkdir(parents=True, exist_ok=True)
    fig.savefig(FIG, dpi=320, bbox_inches="tight")
    plt.close(fig)


def compact_ranking_figure():
    """Plot only the ranking results needed in the conference manuscript."""
    import matplotlib.pyplot as plt
    names = ["DTW + incipit", "Rang sa više\nosobina", "Blaga kazna\nintervala"]
    all_scores = [42.78, 56.11, 53.89]
    mlu_scores = [44.67, 54.67, 56.00]
    pos = range(len(names))
    fig, ax = plt.subplots(figsize=(5.8, 2.5))
    ax.barh([p - .18 for p in pos], all_scores, height=.34, color="#356b88",
            label="Svi automatski kandidati (n=299)")
    ax.barh([p + .18 for p in pos], mlu_scores, height=.34, color="#cf8346",
            label="MLU podskup (n=238)")
    ax.set_yticks(list(pos), names, fontsize=8)
    ax.invert_yaxis()
    ax.set_xlim(0, 65)
    ax.set_xlabel("DA među prvih 20% kandidata za pregled (%)", fontsize=8)
    ax.tick_params(labelsize=8)
    ax.grid(axis="x", alpha=.25, linewidth=.4)
    ax.legend(fontsize=7, loc="lower right")
    fig.tight_layout(pad=.45)
    fig.savefig(RANK_FIG, dpi=320, bbox_inches="tight")
    plt.close(fig)


def para(doc, text="", style=None, align=WD_ALIGN_PARAGRAPH.JUSTIFY):
    p = doc.add_paragraph(style=style)
    p.alignment = align
    p.paragraph_format.space_after = Pt(2)
    p.paragraph_format.line_spacing = 1.0
    p.add_run(text)
    return p


def heading(doc, text):
    p = para(doc, text, "Heading 1", WD_ALIGN_PARAGRAPH.LEFT)
    p.paragraph_format.space_before = Pt(5)
    p.paragraph_format.space_after = Pt(2)
    return p


def subheading(doc, text):
    p = para(doc, text, "Heading 2", WD_ALIGN_PARAGRAPH.LEFT)
    p.paragraph_format.space_before = Pt(3)
    p.paragraph_format.space_after = Pt(1)
    return p


def picture(doc, path, caption):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(2)
    p.paragraph_format.space_after = Pt(1)
    p.add_run().add_picture(str(path), width=Cm(7.55))
    c = para(doc, caption, "Caption", WD_ALIGN_PARAGRAPH.JUSTIFY)
    c.paragraph_format.space_after = Pt(3)


def wide_picture(doc, path, caption, width=15.6):
    before = doc.add_section(WD_SECTION.CONTINUOUS)
    before.page_width, before.page_height = Cm(21), Cm(29.7)
    before.top_margin, before.bottom_margin = Cm(1.8), Cm(1.8)
    before.left_margin, before.right_margin = Cm(1.8), Cm(1.8)
    one_column(before)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(3)
    p.paragraph_format.space_after = Pt(1)
    p.add_run().add_picture(str(path), width=Cm(width))
    c = para(doc, caption, "Caption", WD_ALIGN_PARAGRAPH.JUSTIFY)
    c.paragraph_format.space_after = Pt(3)
    after = doc.add_section(WD_SECTION.CONTINUOUS)
    after.page_width, after.page_height = Cm(21), Cm(29.7)
    after.top_margin, after.bottom_margin = Cm(1.8), Cm(1.8)
    after.left_margin, after.right_margin = Cm(1.8), Cm(1.8)
    two_columns(after)


def build():
    if not FIG.exists():
        compact_example_figure()
    if not RANK_FIG.exists():
        compact_ranking_figure()
    doc = Document()
    first = doc.sections[0]
    first.page_width, first.page_height = Cm(21), Cm(29.7)
    first.top_margin, first.bottom_margin = Cm(1.8), Cm(1.8)
    first.left_margin, first.right_margin = Cm(1.8), Cm(1.8)

    set_font(doc.styles["Normal"], size=9)
    set_font(doc.styles["Title"], size=18, bold=True)
    set_font(doc.styles["Heading 1"], size=10, bold=True)
    set_font(doc.styles["Heading 2"], size=9, bold=True)
    set_font(doc.styles["Caption"], size=7.4)
    doc.styles["Normal"].paragraph_format.space_after = Pt(0)
    doc.styles["Normal"].paragraph_format.line_spacing = 1.0

    header = first.header.paragraphs[0]
    header.text = "IEEESTEC – Studentska konferencija, Niš, 2026."
    header.alignment = WD_ALIGN_PARAGRAPH.LEFT
    for run in header.runs:
        run.font.name = "Times New Roman"
        run.font.size = Pt(8)

    title = doc.add_paragraph("Izdvajanje call-and-response parova iz jazz improvizacija: JazzDialog baza", "Title")
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.paragraph_format.space_after = Pt(5)

    authors = doc.add_table(rows=1, cols=1)
    authors.alignment = WD_TABLE_ALIGNMENT.CENTER
    cell = authors.cell(0, 0)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    p = cell.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(0)
    r = p.add_run("Nikolina Zdravković")
    r.bold = True
    r.font.size = Pt(9)
    p = cell.add_paragraph("Istraživačka stanica Petnica, Valjevo, Srbija")
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(0)
    p.runs[0].font.size = Pt(8)
    p = cell.add_paragraph("Mentorka: Maja Milović")
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(2)
    p.runs[0].font.size = Pt(8)
    for border in cell._tc.tcPr.xpath("./w:tcBorders/*"):
        border.getparent().remove(border)

    body = doc.add_section(WD_SECTION.CONTINUOUS)
    body.page_width, body.page_height = Cm(21), Cm(29.7)
    body.top_margin, body.bottom_margin = Cm(1.8), Cm(1.8)
    body.left_margin, body.right_margin = Cm(1.8), Cm(1.8)
    two_columns(body)

    p = para(doc, "Sažetak — U radu je prikazan postupak izdvajanja call-and-response parova iz transkribovanih jazz solaža i nastanak baze JazzDialog. Kandidati su dobijani iz anotiranih granica fraza, poređeni melodijskim, intervalskim, konturama i ritmičkim osobinama, a zatim rangirani za ručni pregled. Posebna pažnja posvećena je dinamičkom vremenskom poravnanju (DTW). Analiza pokazuje da mala melodijska udaljenost nije dovoljna potvrda call-and-response odnosa: odbijeni kandidat može imati niži DTW skor od prihvaćenog. Konačni postupak zato kombinuje računarsko izdvajanje i rangiranje sa slušanjem i ručnom potvrdom. JazzDialog sadrži 114 potvrđenih zapisa iz 76 solaža, 73 naslova i 42 izvođača, sa preciznim granicama segmenata i MIDI isečcima.")
    p.runs[0].bold = True
    p = para(doc, "Ključne reči — jazz improvizacija, call-and-response, analiza melodije, DTW, MIDI, baza podataka")
    p.runs[0].italic = True

    heading(doc, "I. UVOD")
    para(doc, "Call-and-response je odnos dve uzastopne muzičke celine u kome druga celina odgovara, menja ili razvija ideju prve. U jazz improvizaciji takav odnos može nastati u samoj solaži, bez jasne granice između običnog ponavljanja, varijacije i odgovora. Zbog toga je automatsko izdvajanje zahtevno: slični nizovi nota nisu nužno muzički dijalog, a dobar response ne mora biti melodijski isti kao call.")
    para(doc, "Cilj rada bio je da se ispita kako se iz stvarnih jazz solaža mogu izdvojiti kandidati za call-and-response i da se formira proverljiva baza potvrđenih parova. Istraživačko pitanje je kako numeričke osobine melodijskih segmenata pomažu pri pronalaženju muzički smislenog call-and-response odnosa. Konačni rezultat je JazzDialog, skup ručno potvrđenih parova sa vezom ka izvornim transkripcijama i MIDI isečcima.")

    heading(doc, "II. METOD")
    subheading(doc, "A. Podaci i kandidati")
    para(doc, "Korišćene su monofone MIDI transkripcije solo improvizacija i prateće anotacije sistema Jazzomat / Weimar Jazz Database [1]. Svaka nota sadrži redni broj, početak, MIDI visinu i trajanje. Zvanično anotirane granice fraza korišćene su kao stabilan početni okvir, jer se ranija segmentacija pauzama pokazala previše krutom, a klizni prozori kroz ceo solo davali su previše lokalnih, muzički nepovezanih podudaranja.")
    para(doc, "U osnovnoj pretrazi jedna fraza se deli na dva uzastopna dela. Za svaku dozvoljenu tačku podele prvi deo je call, a drugi response. Kandidati sa premalo nota, prekratkim trajanjem ili izrazito neujednačenim dužinama su odbačeni. U kasnijoj verziji response je imao između polovine i dvostruke dužine call-a, a oba dela su trajala najmanje jednu sekundu. Ta ograničenja ne definišu call-and-response, već uklanjaju trivijalne slučajeve.")

    subheading(doc, "B. Numeričke osobine i rangiranje")
    para(doc, "Osnovna melodijska mera bilo je dinamičko vremensko poravnanje (DTW), koje poredi dve sekvence i dozvoljava lokalno poravnanje segmenata nejednake dužine [2]. Korišćen je optimalni put P kroz matricu razlika MIDI visina; rastojanje je koren zbira kvadrata razlika duž puta, podeljen njegovom stvarnom dužinom. Manja vrednost označava bliže melodijsko poravnanje.")
    para(doc, "Pored apsolutnih MIDI visina korišćena je transponovana, odnosno shape predstava, dobijena oduzimanjem prve note segmenta. Ispitan je incipit, poređenje prvih nota call-a i response-a; u vrlo brzim frazama početak je proširivan do približno 0,75 s. Razmatrani su i intervalski nizovi, sažeta melodijska kontura, relativni ritam zasnovan na međunotnim razmacima, median filter i ograničeno poravnanje.")
    para(doc, "U kasnijem koraku kandidati su rangirani korišćenjem DTW-a, incipita, odnosa dužina, zastupljenosti istih tonova, raznovrsnosti visina i približnog ponavljanja intervala. Negativno označeni primeri korišćeni su samo za blage kazne čestih loših obrazaca; obrazac nije strogo zabranjen, jer ponavljanje može biti deo validnog odgovora.")

    subheading(doc, "C. Potvrda")
    para(doc, "Svaki predlog je izvezen u CSV tabelu i MIDI isečak sa razdvojenim call i response delom. Nakon slušanja je ručno označen kao DA ili NE. Sistem zato automatski pronalazi i rangira kandidate, dok muzički odnos potvrđuje slušanje. Oznake je obavila jedna anotatorka.")

    heading(doc, "III. REZULTATI I DISKUSIJA")
    para(doc, "Najjasniji nalaz odnosi se na ograničenje melodijske sličnosti. Slika 1 poredi dva stvarno pregledana kandidata. Prihvaćeni par iz sola „I Fall in Love Too Easily“ ima DTW skor 0,464, dok odbijeni par iz „Blues for Blanche“ ima manji skor, 0,125. U drugom slučaju kratko, mehaničko ponavljanje tonova omogućava jeftino DTW poravnanje, ali se pri slušanju ne doživljava kao odgovor na prethodnu ideju. U prvom slučaju response nije kopija call-a, već razvija srodnu melodijsku misao.")
    wide_picture(doc, FIG, "Slika 1. Melodijske konture call-a i response-a za (a) prihvaćeni i (b) odbijeni kandidat. Horizontalna osa je redni broj note unutar segmenta, a vertikalna osa MIDI visina. Odbijeni primer ima znatno niži DTW skor, što pokazuje da skor sam nije odluka o muzičkoj funkciji.")
    para(doc, "Nalaz se vidi i na većem razvojnom skupu od 299 automatski izdvojenih kandidata: 101 je označen sa DA, a 198 sa NE. Na slici 2 horizontalna osa razdvaja dve grupe, dok vertikalna prikazuje globalni pitch-DTW skor podeljen dužinom puta. Violine pokazuju raspodelu, pojedinačne tačke kandidate, a crna linija centralni položaj raspodele. Grupe se snažno preklapaju, pa jedan prag može smanjiti broj predloga, ali ne može pouzdano potvrditi call-and-response vezu.")
    wide_picture(doc, OUT / "figure" / "figure_2_dtw_overlap.png", "Slika 2. Raspodele normalizovanog globalnog DTW skora za ručno prihvaćene i odbijene automatske kandidate.", width=15.6)
    para(doc, "Na istom zamrznutom razvojnom skupu testirano je rangiranje kandidata uz petostruku podelu po solima i tri semena. Mera nije accuracy klasifikatora: predstavlja procenat DA kandidata među prvih 20% rangirane liste za slušanje. DTW sa incipitom davao je prosečno 42,78%, raniji rang sa više karakteristika 56,11%, a rang sa blagom kaznom intervalskih obrazaca 53,89%. Dakle, više osobina može poboljšati redosled pregleda, ali rezultat i dalje nije dovoljan da zameni slušanje.")
    wide_picture(doc, RANK_FIG, "Slika 3. Udeo ručno prihvaćenih kandidata u prvih 20% rangirane liste. Stubovi upoređuju sve automatske kandidate i MLU podskup; prikaz je sažet na tri najinformativnija postupka.", width=15.6)
    para(doc, "Pojedinačne mere ritma i konture nisu dale stabilno razdvajanje, a samostalna učestalost intervalskih obrazaca dala je 35,00%. To pokazuje da obrazac koji je čest među odbijenim primerima ne treba strogo zabraniti: ponavljanje je ponekad upravo deo validnog odgovora.")
    para(doc, "Rani skup od 65 ručno označenih kandidata iz 22 solaže korišćen je za proveru više melodijskih predstava. AUC je bio 0,475 za početnu DTW meru, 0,567 za pitch RMS, 0,557 za transponovani shape RMS i 0,603 za shape RMS sa ograničenom dijagonalnom trakom. Poslednja varijanta je bila najbolja u tom poređenju, ali nije bila dovoljna za autonomno izdvajanje. Prosta ritmička mera zasnovana na relativnim međunotnim razmacima imala je AUC 0,501, bez korisnog razdvajanja.")
    para(doc, "Kontura je zato posmatrana kao dodatna, a ne kao samostalna odluka. Ideju je podstakla statistička karakterizacija melodijskih kontura [4], ali se postupak ne može direktno preneti: taj rad izdvaja vodeću melodiju iz polifonog zvuka, dok su ovde već dostupne monofone MIDI note i treba proceniti odnos dva segmenta. Sažeta kontura i njena kombinacija sa ritmom nisu dale stabilan napredak.")
    para(doc, "Konačni skup JazzDialog sadrži 114 ručno potvrđenih call-and-response zapisa iz 76 solaža, 73 naslova i 42 izvođača. Svaki zapis čuva identifikator sola, izvođača, naslov, granice i vremena call/response delova, nizove MIDI visina i putanju do MIDI isečka. Call delovi imaju medijanu od 9, a response delovi medijanu od 10 nota. Dvadeset šest anotacija u celosti je unutar jedne zvanične fraze, dok 88 prelazi njenu granicu.")
    para(doc, "Za razliku od skupova namenjenih generisanju muzičkog odgovora [3], JazzDialog polazi od stvarnih improvizacija i dokumentuje postupak pronalaženja i ručne potvrde odnosa koji se već dogodio. Njegov doprinos je proverljiva kolekcija i negativan rezultat važan za dalji rad: melodijska sličnost hvata deo odnosa, ali ne obuhvata njegovu muzičku funkciju.")

    heading(doc, "IV. OGRANIČENJA I DALJI RAD")
    para(doc, "Call-and-response je delimično subjektivan pojam. Granice između doslovnog ponavljanja, varijacije i odgovora nisu uvek iste za različite slušaoce, a u ovoj verziji svaku konačnu oznaku dala je jedna anotatorka. Zbog toga rezultati opisuju transparentan postupak formiranja baze, ali ne predstavljaju meru saglasnosti muzičara. Sledeća verzija treba da uključi nezavisne oznake više slušalaca i analizu međuanotatorske saglasnosti.")
    para(doc, "Korišćene su monofone MIDI transkripcije. One čuvaju visinu, početak i trajanje note, ali ne obuhvataju artikulaciju, dinamiku, boju tona, harmonijsku pratnju ni interakciju sa ansamblom. Negativan rezultat jednostavne ritmičke mere zato ne znači da ritam nema perceptivnu ulogu, već da poređenje međunotnih razmaka nije bilo dovoljno za ovu odluku.")
    para(doc, "Razvojne oznake su korišćene i pri oblikovanju kriterijuma i pri retrospektivnom poređenju rangiranja. Petostruka podela po solima smanjuje curenje srodnih primera između učenja i provere, ali nije potpuno nezavisna završna evaluacija. Budući rad treba unapred da izdvoji i zamrzne skup solaža za slepu anotaciju i testiranje.")
    para(doc, "JazzDialog se može proširivati istim formatom, sa preciznim granicama i MIDI isečcima. Veća baza bi omogućila pouzdanije poređenje melodijskih, ritmičkih i harmonskih osobina, kao i modele koji bi davali prioritet kandidatima za pregled. Takav model bi i dalje bio pomoć pri anotiranju, a ne zamena za muzičko slušanje.")

    heading(doc, "V. ZAKLJUČAK")
    para(doc, "Računarska analiza može efikasno izdvojiti i poređati kandidate za call-and-response u jazz solažama, ali nijedna ispitana pojedinačna mera nije bila dovoljna za pouzdanu konačnu odluku. DTW je bio koristan za poređenje i vizuelnu analizu, ali primeri i raspodele pokazuju da manja udaljenost ne znači nužno bolji muzički odgovor. Konačni hibridni postupak — izdvajanje i rangiranje kandidata, preslušavanje i ručna potvrda — doveo je do baze JazzDialog sa 114 parova. Ograničenja su subjektivnost pojma, jedna anotatorka, monofone transkripcije i činjenica da su razvojne oznake korišćene pri oblikovanju kriterijuma. Sledeći koraci su nezavisna anotacija više muzičara, proširenje baze i učenje na većem skupu.")

    heading(doc, "LITERATURA")
    references = [
        "[1] M. Pfleiderer, K. Frieler, J. Abeßer, W. G. Zaddach i B. Burkhart, ur., Inside the Jazzomat: New Perspectives for Jazz Research. Mainz: Schott Campus, 2017.",
        "[2] M. Müller, Fundamentals of Music Processing: Audio, Analysis, Algorithms, Applications. Cham: Springer, 2015, doi: 10.1007/978-3-319-21945-5.",
        "[3] Y. Hu, Z. Wang, R. Liu, Y. Liang i Y. Zhang, “Responding to the Call: Exploring Automatic Music Composition Using a Knowledge-Enhanced Model,” Proc. AAAI Conf. Artif. Intell., vol. 38, no. 1, pp. 521–529, 2024, doi: 10.1609/aaai.v38i1.27807.",
        "[4] J. Salamon, G. Peeters i A. Röbel, “Statistical Characterisation of Melodic Pitch Contours and Its Application for Melody Extraction,” Proc. 13th Int. Soc. Music Information Retrieval Conf., pp. 187–192, 2012.",
    ]
    for ref in references:
        p = para(doc, ref, align=WD_ALIGN_PARAGRAPH.LEFT)
        p.paragraph_format.left_indent = Cm(.18)
        p.paragraph_format.first_line_indent = Cm(-.18)

    doc.core_properties.title = "Izdvajanje call-and-response parova iz jazz improvizacija: JazzDialog baza"
    doc.core_properties.author = "Nikolina Zdravković"
    doc.core_properties.subject = "IEEESTEC konferencijski rad"
    output = OUT / "ieeestec_jazzdialog.docx"
    doc.save(output)
    print(output)


if __name__ == "__main__":
    build()
